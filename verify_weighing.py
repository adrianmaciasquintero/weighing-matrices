#!/usr/bin/env python3
"""
Exhaustively verifies that each solution_*.json in the solutions folder
contains a weighing matrix W(n, k).

A matrix W of order n and weight k is a weighing matrix if:
  1. It is square n x n with entries in {-1, 0, +1}.
  2. Each row has exactly k nonzero entries.
  3. Each column has exactly k nonzero entries.
  4. W @ W.T == k * I   (orthogonal rows, each with norm^2 = k).
  5. W.T @ W == k * I   (orthogonal columns, each with norm^2 = k).

All matrix checks use int64 arithmetic and np.array_equal. There is NO
floating point, NO tolerance, NO allclose anywhere in the matrix verdict.
For entries in {-1,0,1} and n <= ~1e9 the products never overflow int64.

JSON consistency checks (warnings; they do NOT change the matrix verdict):
  - matrix_symbols matches matrix  ('+' -> 1, '-' -> -1, '.' -> 0)
  - n == 4*m
  - folder name W{n}_{k} consistent with n and k
  - method-specific checks: gs, identity, h2blocks, dsum, kron, paley*

Usage:
    python verify_weighing.py
    python verify_weighing.py /path/to/solutions
    python verify_weighing.py /path/to/solutions --first-only
    python verify_weighing.py /path/to/solutions --verbose
    python verify_weighing.py /path/to/solutions --export report.csv
    python verify_weighing.py /path/to/solutions --export report.xlsx
"""

__author__ = "Adrian Macias-Quintero"
__version__ = "1.1.0"
__license__ = "Apache-2.0"
__copyright__ = "Copyright (c) 2025 Adrian Macias-Quintero"

import argparse
import csv
import json
import re
import sys
from pathlib import Path

import numpy as np

DEFAULT_PATH = "solutions"
REQUIRED_KEYS = ["n", "k", "m", "matrix"]
EXPORT_FIELDS = [
    "n", "k", "folder", "file", "method",
    "status", "num_errors", "num_warnings", "errors", "warnings",
]


# ----------------------------------------------------------------------
# Goethals-Seidel construction
# ----------------------------------------------------------------------
def circulant(x):
    m = len(x)
    return np.array([[x[(j - i) % m] for j in range(m)] for i in range(m)], dtype=np.int64)


def goethals_seidel(A, B, C, D):
    m = len(A)
    a, b, c, d = (circulant(v) for v in (A, B, C, D))
    R = np.fliplr(np.eye(m, dtype=np.int64))
    blocks = [
        [a,        b @ R,    c @ R,     d @ R],
        [-b @ R,   a,        d.T @ R,  -c.T @ R],
        [-c @ R,  -d.T @ R,  a,         b.T @ R],
        [-d @ R,   c.T @ R, -b.T @ R,   a],
    ]
    return np.block(blocks)


def get_method(data):
    """Return the method name, from 'method' or 'recipe.type'."""
    if isinstance(data.get("method"), str):
        return data["method"]
    r = data.get("recipe")
    if isinstance(r, dict) and isinstance(r.get("type"), str):
        return r["type"]
    if data.get("A"):  # old files without 'method' but with GS generators
        return "gs"
    return ""


# ----------------------------------------------------------------------
# Matrix verification  (pure int64, exact)
# ----------------------------------------------------------------------
def verify(data, folder_name=""):
    """Returns (errors, warnings). Errors => NOT a valid weighing matrix."""
    errors, warnings = [], []

    missing = [key for key in REQUIRED_KEYS if key not in data]
    if missing:
        errors.append(f"missing keys in JSON: {missing}")
        return errors, warnings

    n, k = data["n"], data["k"]
    if not all(isinstance(v, int) and not isinstance(v, bool) for v in (n, k)):
        errors.append("n, k must be integers")
        return errors, warnings

    # --- 1. shape and values ---
    try:
        W = np.array(data["matrix"], dtype=np.int64)
    except Exception as e:
        errors.append(f"matrix is not a rectangular numeric matrix: {e}")
        return errors, warnings

    if W.ndim != 2 or W.shape != (n, n):
        errors.append(f"dimension {W.shape} != ({n}, {n})")
        return errors, warnings

    values = set(np.unique(W).tolist())
    if not values <= {-1, 0, 1}:
        errors.append(f"entries outside {{-1,0,1}}: {sorted(values - {-1, 0, 1})}")
        return errors, warnings

    # --- 2 and 3. number of nonzero entries per row and column ---
    nonzero_rows = np.count_nonzero(W, axis=1)
    nonzero_columns = np.count_nonzero(W, axis=0)
    bad_rows = np.where(nonzero_rows != k)[0]
    bad_columns = np.where(nonzero_columns != k)[0]
    if len(bad_rows):
        errors.append(f"{len(bad_rows)} row(s) with weight != {k} "
                      f"(e.g. row {bad_rows[0]} has {nonzero_rows[bad_rows[0]]})")
    if len(bad_columns):
        errors.append(f"{len(bad_columns)} column(s) with weight != {k} "
                      f"(e.g. col {bad_columns[0]} has {nonzero_columns[bad_columns[0]]})")

    # --- 4. row orthogonality (exact int64) ---
    target = k * np.eye(n, dtype=np.int64)
    G = W @ W.T                    # int64 @ int64 -> int64, no BLAS, no float
    if not np.array_equal(G, target):
        bad_diag = int(np.count_nonzero(np.diag(G) != k))
        off = G - np.diag(np.diag(G))
        bad_off = int(np.count_nonzero(off))
        errors.append(f"W·W^T != {k}·I (incorrect diagonal: {bad_diag}, "
                      f"nonzero products outside the diagonal: {bad_off})")

    # --- 5. column orthogonality (exact int64) ---
    G2 = W.T @ W
    if not np.array_equal(G2, target):
        bad_diag = int(np.count_nonzero(np.diag(G2) != k))
        off = G2 - np.diag(np.diag(G2))
        bad_off = int(np.count_nonzero(off))
        errors.append(f"W^T·W != {k}·I (incorrect diagonal: {bad_diag}, "
                      f"nonzero products outside the diagonal: {bad_off})")

    # ------------------------------------------------------------------
    # JSON consistency checks (warnings only)
    # ------------------------------------------------------------------
    method = get_method(data)

    symbols = data.get("matrix_symbols")
    if symbols is not None:
        expected = ["".join("+" if v == 1 else "-" if v == -1 else "." for v in row)
                    for row in W]
        if list(symbols) != expected:
            warnings.append("matrix_symbols does not match matrix")

    m = data.get("m")
    if isinstance(m, int) and n != 4 * m:
        warnings.append(f"n={n} != 4*m={4 * m}")

    match = re.match(r"W(\d+)_(\d+)", folder_name)
    if match:
        folder_n, folder_k = int(match.group(1)), int(match.group(2))
        if folder_n != n:
            warnings.append(f"folder '{folder_name}' suggests n={folder_n} but JSON has n={n}")
        if folder_k != k:
            warnings.append(f"folder '{folder_name}' suggests k={folder_k} but JSON has k={k}")

    recipe = data.get("recipe") if isinstance(data.get("recipe"), dict) else None

    if method == "gs" or (not recipe and data.get("A")):
        _check_gs(data, W, n, k, m, warnings)
    elif method == "identity":
        if not np.array_equal(W, np.eye(n, dtype=np.int64)):
            warnings.append("method is 'identity' but matrix is not the identity")
    elif method == "h2blocks":
        H2 = np.array([[1, 1], [1, -1]], dtype=np.int64)
        expected = np.kron(np.eye(n // 2, dtype=np.int64), H2)
        if not np.array_equal(W, expected):
            warnings.append("method is 'h2blocks' but matrix != kron(I, H2)")
    elif method == "dsum":
        _check_dsum(data, W, n, k, recipe, warnings)
    elif method == "kron":
        _check_kron(data, W, n, k, recipe, warnings)
    elif method in ("paley1", "paley1_h", "paley2"):
        _check_paley(method, W, n, k, recipe, warnings)
    elif not method:
        warnings.append("could not determine 'method' (missing 'method' and 'recipe.type')")

    return errors, warnings


# ----------------------------------------------------------------------
# Method-specific consistency helpers
# ----------------------------------------------------------------------
def _check_gs(data, W, n, k, m, warnings):
    A, B, C, D = (data.get(x) for x in "ABCD")
    if A is None or B is None or C is None or D is None:
        return
    sequences = [A, B, C, D]
    if any(len(seq) != m for seq in sequences):
        warnings.append(f"some sequence A,B,C,D does not have length m={m}")
        return
    if any(abs(v) > 1 for seq in sequences for v in seq):
        warnings.append("A,B,C,D have entries outside {-1,0,1}")

    weights = [int(np.count_nonzero(seq)) for seq in sequences]
    sums = [int(sum(seq)) for seq in sequences]
    if list(data.get("weights", [])) != weights:
        warnings.append(f"weights {data.get('weights')} != actual nonzero counts {weights}")
    if sum(weights) != k:
        warnings.append(f"sum(weights)={sum(weights)} != k={k}")
    if list(data.get("sums", [])) != sums:
        warnings.append(f"sums {data.get('sums')} != actual sums {sums}")
    if sum(s * s for s in sums) != k:
        warnings.append(f"sum(sums^2)={sum(s * s for s in sums)} != k={k}")

    if n == 4 * m:
        GS = goethals_seidel(*sequences)
        if not np.array_equal(GS, W):
            warnings.append("matrix does not match the Goethals-Seidel array of A,B,C,D "
                            "(it may use another convention)")


def _check_dsum(data, W, n, k, recipe, warnings):
    if recipe is None or recipe.get("type") != "dsum":
        warnings.append("method is 'dsum' but recipe is missing or does not match")
        return
    parts = recipe.get("parts")
    if not (isinstance(parts, list) and len(parts) == 2
            and all(isinstance(p, (list, tuple)) and len(p) == 2 for p in parts)):
        warnings.append(f"recipe['parts'] malformed: {parts!r}")
        return
    (a, ka), (b, kb) = parts
    if a + b != n:
        warnings.append(f"dsum parts sizes {a}+{b} != n={n}")
    if ka != k or kb != k:
        warnings.append(f"dsum parts weights ({ka},{kb}) != k={k}")
    if 0 < a < n and 0 < b < n:
        off = W[:a, a:]
        off2 = W[a:, :a]
        if off.any() or off2.any():
            warnings.append("method is 'dsum' but matrix has nonzero entries "
                            "outside the block-diagonal")


def _check_kron(data, W, n, k, recipe, warnings):
    if recipe is None or recipe.get("type") != "kron":
        warnings.append("method is 'kron' but recipe is missing or does not match")
        return
    left = recipe.get("left")
    right = recipe.get("right")
    if not (isinstance(left, (list, tuple)) and len(left) == 2
            and isinstance(right, (list, tuple)) and len(right) == 2):
        warnings.append(f"kron recipe malformed: left={left!r}, right={right!r}")
        return
    a, k1 = left
    b, k2 = right
    if a * b != n:
        warnings.append(f"kron sizes {a}*{b} != n={n}")
    if k1 * k2 != k:
        warnings.append(f"kron weights {k1}*{k2} != k={k}")


def _check_paley(method, W, n, k, recipe, warnings):
    if recipe is None or recipe.get("type") != method:
        warnings.append(f"method is '{method}' but recipe is missing or does not match")
        return
    q = recipe.get("q")
    if not isinstance(q, int):
        warnings.append(f"{method} recipe has no integer 'q'")
        return
    if method == "paley1":
        if n != q + 1 or k != q:
            warnings.append(f"paley1 expected n=q+1={q + 1}, k=q={q}; got n={n}, k={k}")
    elif method == "paley1_h":
        if n != q + 1 or k != q + 1:
            warnings.append(f"paley1_h expected n=q+1={q + 1}, k=q+1={q + 1}; got n={n}, k={k}")
    elif method == "paley2":
        if n != 2 * (q + 1) or k != 2 * (q + 1):
            warnings.append(f"paley2 expected n=k=2(q+1)={2 * (q + 1)}; got n={n}, k={k}")


# ----------------------------------------------------------------------
# Export
# ----------------------------------------------------------------------
def _join_messages(msgs):
    return " | ".join(msgs) if msgs else ""


def export_rows(rows, path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        _export_csv(rows, p)
    elif ext in (".xlsx", ".xls"):
        _export_xlsx(rows, p)
    else:
        raise ValueError(f"unsupported export extension '{ext}' (use .csv or .xlsx)")
    print(f"\nReport written to: {p.resolve()}")


def _export_csv(rows, path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=EXPORT_FIELDS)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def _export_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("\n[warn] openpyxl is not installed; cannot write .xlsx. "
              "Falling back to CSV with the same base name.")
        _export_csv(rows, path.with_suffix(".csv"))
        print(f"Report written to: {path.with_suffix('.csv').resolve()}")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "verification"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    fail_fill = PatternFill("solid", fgColor="FCE4E4")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    ws.append(EXPORT_FIELDS)
    for c in range(1, len(EXPORT_FIELDS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([r[field] for field in EXPORT_FIELDS])

    for i, r in enumerate(rows, start=2):
        status = r["status"]
        fill = fail_fill if status == "FAIL" else (warn_fill if r["num_warnings"] else ok_fill)
        for c in range(1, len(EXPORT_FIELDS) + 1):
            ws.cell(row=i, column=c).fill = fill

    widths = {"n": 6, "k": 6, "folder": 14, "file": 18, "method": 16,
              "status": 8, "num_errors": 11, "num_warnings": 13,
              "errors": 60, "warnings": 60}
    for idx, field in enumerate(EXPORT_FIELDS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(field, 14)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)


# ----------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Verifies weighing matrices W(n,k) in solutions/")
    parser.add_argument("path", nargs="?", default=DEFAULT_PATH)
    parser.add_argument("--first-only", action="store_true",
                        help="only solution_001.json from each folder")
    parser.add_argument("--verbose", action="store_true",
                        help="also show consistency warnings")
    parser.add_argument("--export", default=None,
                        help="export the report to a .csv or .xlsx file")
    args = parser.parse_args()

    base = Path(args.path)
    pattern = "W*_*/solution_001.json" if args.first_only else "W*_*/solution_*.json"
    paths = sorted(base.glob(pattern))
    print(f"Files found: {len(paths)}\n")
    if not paths:
        sys.exit(1)

    ok = failed = with_warnings = 0
    rows = []
    for path in paths:
        label = f"{path.parent.name}/{path.name}"
        data = {}
        try:
            with open(path, "r", encoding="utf-8") as file:
                data = json.load(file)
            errors, warnings = verify(data, path.parent.name)
        except Exception as e:
            errors, warnings = [f"could not read/verify: {e}"], []

        n = data.get("n", "")
        k = data.get("k", "")
        method = get_method(data)

        if errors:
            failed += 1
            print(f"[FAIL]  {label}")
            for error in errors:
                print(f"        - {error}")
            status = "FAIL"
        else:
            ok += 1
            marker = " (with warnings)" if warnings else ""
            extra = f" [{method}]" if method else ""
            print(f"[OK]    {label}  ->  W({n},{k}) valid{extra}{marker}")
            status = "OK"
        if warnings:
            with_warnings += 1
            if args.verbose or errors:
                for warning in warnings:
                    print(f"        ! warning: {warning}")

        rows.append({
            "n": n,
            "k": k,
            "folder": path.parent.name,
            "file": path.name,
            "method": method,
            "status": status,
            "num_errors": len(errors),
            "num_warnings": len(warnings),
            "errors": _join_messages(errors),
            "warnings": _join_messages(warnings),
        })

    print("\n" + "=" * 60)
    print(f"Total: {len(paths)} | OK: {ok} | FAIL: {failed} | "
          f"with consistency warnings: {with_warnings}")
    print("Author: Adrian Macias-Quintero")
    if with_warnings and not args.verbose:
        print("(use --verbose to see the warnings)")

    if args.export:
        export_rows(rows, args.export)

    sys.exit(0 if failed == 0 else 1)


if __name__ == "__main__":
    main()